from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "WhatsDrink-产品介绍.docx"
HOME_IMAGE = ROOT / "docs" / "design" / "home-selected.png"
PROFILE_IMAGE = ROOT / "docs" / "design" / "profile-selected.png"

# Design preset: standard_business_brief.
# Named WhatsDrink brand override: coffee brown #4A2E1F, coral #D96C4A,
# cream #F8F1E5, muted brown #8B735F. CJK font override: Heiti SC.
COFFEE = "4A2E1F"
CORAL = "D96C4A"
CREAM = "F8F1E5"
MUTED = "8B735F"
LIGHT = "F3E8D8"
GRID = "DCCDBB"
WHITE = "FFFFFF"
TOTAL_WIDTH = 9360


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def apply_cjk_font(r_pr, name="Heiti SC"):
    r_fonts = r_pr.rFonts
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), name)
    r_fonts.set(qn("w:cs"), name)
    r_fonts.set(qn("w:hint"), "eastAsia")
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:val"), "zh-CN")
    lang.set(qn("w:eastAsia"), "zh-CN")


def set_run_font(run, size=None, bold=None, color=None, name="Heiti SC"):
    run.font.name = name
    apply_cjk_font(run._element.get_or_add_rPr(), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_run_font(run, **kwargs)
    return run


def set_paragraph_shading(paragraph, fill, border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if border:
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "24")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border)
        borders.append(left)
        p_pr.append(borders)


def set_paragraph_keep(paragraph, keep_next=False, keep_lines=True):
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        p_pr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        p_pr.append(OxmlElement("w:keepLines"))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(paragraph, "WhatsDrink  ·  ", size=9, color=MUTED)
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_separate, fld_end])
    set_run_font(run, size=9, color=MUTED)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    add_text(p, text, bold=True, color=COFFEE)
    set_paragraph_keep(p, keep_next=True)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Body Text")
    if bold_lead and text.startswith(bold_lead):
        add_text(p, bold_lead, bold=True, color=COFFEE)
        add_text(p, text[len(bold_lead):])
    else:
        add_text(p, text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    add_text(p, text)
    return p


def create_numbering_instance(doc):
    numbering = doc.part.numbering_part.element
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering.findall(qn("w:num"))
    ]
    num_id = max(num_ids, default=0) + 1

    style = doc.styles["List Number"]._element
    base_num_node = style.find(f".//{qn('w:numId')}")
    base_num_id = int(base_num_node.get(qn("w:val")))
    base_num = next(
        node
        for node in numbering.findall(qn("w:num"))
        if int(node.get(qn("w:numId"))) == base_num_id
    )
    abstract_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), abstract_id)
    num.append(abstract_ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return num_id


def add_number(doc, lead, text, num_id):
    p = doc.add_paragraph(style="List Number")
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_ref = OxmlElement("w:numId")
    num_ref.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_ref])
    p._p.get_or_add_pPr().append(num_pr)
    p.paragraph_format.space_after = Pt(8)
    add_text(p, lead, bold=True, color=COFFEE)
    add_text(p, text)
    return p


def add_callout(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.15)
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(12)
    set_paragraph_shading(p, CREAM, CORAL)
    add_text(p, text, size=12, bold=True, color=COFFEE)
    set_paragraph_keep(p)
    return p


def add_page_break(doc):
    doc.add_page_break()


def add_image_with_alt(paragraph, path, width, alt):
    run = paragraph.add_run()
    inline = run.add_picture(str(path), width=Inches(width))
    inline._inline.docPr.set("descr", alt)
    inline._inline.docPr.set("title", alt)
    return inline


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Heiti SC"
    apply_cjk_font(normal._element.get_or_add_rPr())
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(COFFEE)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    body = styles["Body Text"]
    body.base_style = normal
    body.font.name = "Heiti SC"
    apply_cjk_font(body._element.get_or_add_rPr())
    body.font.size = Pt(11)
    body.paragraph_format.space_before = Pt(0)
    body.paragraph_format.space_after = Pt(7)
    body.paragraph_format.line_spacing = 1.10

    for name, size, before, after, color in (
        ("Heading 1", 16, 16, 8, COFFEE),
        ("Heading 2", 13, 12, 6, COFFEE),
        ("Heading 3", 12, 8, 4, MUTED),
    ):
        style = styles[name]
        style.font.name = "Heiti SC"
        apply_cjk_font(style._element.get_or_add_rPr())
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Heiti SC"
        apply_cjk_font(style._element.get_or_add_rPr())
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167

    if "Caption WhatsDrink" not in [s.name for s in styles]:
        caption = styles.add_style("Caption WhatsDrink", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Caption WhatsDrink"]
    caption.font.name = "Heiti SC"
    apply_cjk_font(caption._element.get_or_add_rPr())
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(hp, "WhatsDrink  /  产品介绍", size=9, bold=True, color=MUTED)

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def build():
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])

    # Editorial cover template.
    cover = doc.add_paragraph()
    cover.paragraph_format.space_before = Pt(104)
    cover.paragraph_format.space_after = Pt(20)
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(cover, "W H A T S D R I N K", size=10, bold=True, color=CORAL)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    add_text(title, "WhatsDrink 微信小程序", size=29, bold=True, color=COFFEE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    add_text(subtitle, "记录每一杯，也轻松决定下一杯", size=16, color=MUTED)

    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(rule, "●", size=10, color=CORAL)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(72)
    meta.paragraph_format.space_after = Pt(8)
    add_text(meta, "产品介绍  V1.0", size=11, bold=True, color=COFFEE)

    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(date, "2026 年 7 月", size=10, color=MUTED)

    prepared = doc.add_paragraph()
    prepared.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prepared.paragraph_format.space_before = Pt(42)
    add_text(prepared, "面向产品评审、团队协作与项目演示", size=9, color=MUTED)

    add_page_break(doc)

    add_heading(doc, "产品概览", 1)
    add_callout(
        doc,
        "WhatsDrink 是一款面向日常饮品记录者的个人微信小程序，把“记下今天喝了什么”和“帮我决定下一杯喝什么”连接成一段简单、温暖的日常体验。",
    )
    facts = doc.add_paragraph()
    facts.alignment = WD_ALIGN_PARAGRAPH.CENTER
    facts.paragraph_format.space_after = Pt(14)
    add_text(facts, "2 项核心任务   ·   4 个主入口   ·   1 个私密数据空间", size=11, bold=True, color=CORAL)

    add_heading(doc, "产品背景", 2)
    add_body(
        doc,
        "对于习惯记录生活的人，一杯咖啡或奶茶往往包含了时间、场景和情绪。但这些信息通常散落在备忘录、相册或订单中，难以形成连续的个人记录。",
    )
    add_body(
        doc,
        "与此同时，面对越来越多的品牌和单品，“今天喝什么”也成为一个高频但琐碎的选择题。传统随机工具可以给出结果，却无法继续完成规格选择、热量补充和记录保存。",
    )
    add_body(doc, "WhatsDrink 因此聚焦两个相邻任务：")
    product_task_numbering = create_numbering_instance(doc)
    add_number(
        doc,
        "快速记录。",
        "用尽可能轻的操作，完整记下一杯饮品。",
        product_task_numbering,
    )
    add_number(
        doc,
        "轻松决定。",
        "在犹豫时给出一个明确选择，并自然进入记录流程。",
        product_task_numbering,
    )

    add_heading(doc, "产品定位", 2)
    add_bullet(doc, "目标用户：喜欢记录日常的咖啡与奶茶爱好者；关注饮品热量、价格或口味的人；容易在饮品选择上纠结的人。")
    add_bullet(doc, "核心价值：既能留下喝过什么，也能帮助决定接下来喝什么。")
    add_bullet(doc, "产品角色：私密、轻量、有陪伴感的个人饮品日记。")
    add_bullet(doc, "设计原则：快速记录、诚实呈现、温和陪伴、选择与记录闭环。")

    add_page_break(doc)

    add_heading(doc, "核心功能", 1)
    add_body(doc, "四个主入口覆盖从即时记录、日历回看、选择决策到个人管理的完整使用链路。")
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [1800, 2600, 4960])
    mark_header_row(table.rows[0])
    headers = ["模块", "用户任务", "核心体验"]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT)
        p = cell.paragraphs[0]
        add_text(p, text, size=10, bold=True, color=COFFEE)
    data = [
        ("今日记录", "记下当天喝过的饮品", "时间线展示、已知热量合计、快速新增、编辑与删除"),
        ("饮品记录", "完整描述一杯饮品", "品牌与单品搜索或自定义；规格、温度、甜度、价格、评分、备注和照片"),
        ("记录日历", "回看某一天或某个月", "标记有记录的日期，按日查看杯数、明细与已知热量"),
        ("Choice One", "决定“今天喝什么”", "默认咖啡转盘、自定义多个转盘、等概率抽取、结果预填到新增记录"),
        ("我的", "管理个人资料与数据", "一周饮品日历、资料编辑、隐私说明、清除记录与注销账号"),
    ]
    for row_data in data:
        cells = table.add_row().cells
        for idx, text in enumerate(row_data):
            p = cells[idx].paragraphs[0]
            add_text(p, text, size=9.5, bold=(idx == 0), color=COFFEE if idx == 0 else None)
            set_cell_margins(cells[idx])
            set_cell_width(cells[idx], [1800, 2600, 4960][idx])
            cells[idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    add_heading(doc, "关键体验", 2)
    add_heading(doc, "一杯饮品，快速记下", 3)
    add_body(
        doc,
        "用户可以从已收录的品牌与单品中选择，也可以直接输入自定义饮品。表单支持日期时间、分类、品牌、单品、杯型、温度、甜度、参考热量、价格、评分、备注与一张照片。",
    )
    add_body(
        doc,
        "预置热量可以修改；未知热量不会被错误计入汇总，并会明确提示数据不完整。热量统一标注为参考值，不提供医疗、营养或减脂建议。",
    )
    add_heading(doc, "Choice One，让选择自然进入行动", 3)
    add_body(
        doc,
        "首次使用自动创建“今天喝什么咖啡”默认转盘。用户也可以管理多个转盘，每个转盘容纳 2–20 个饮品库或自由文本候选项。抽取采用等概率机制，结果在动画开始前确定。",
    )
    add_body(
        doc,
        "点击“就喝这个”后，系统打开新增记录并预填饮品信息，但不会替用户自动保存，保留最终确认权。",
    )

    add_page_break(doc)

    add_heading(doc, "典型使用场景", 1)
    scenario_numbering = create_numbering_instance(doc)
    add_number(
        doc,
        "上班路上的快速记录",
        "：用户买到一杯熟悉的拿铁，从饮品库选择品牌和单品，确认规格与参考热量，补充评分后保存。",
        scenario_numbering,
    )
    add_number(
        doc,
        "下午茶前的选择困难",
        "：用户打开 Choice One，在自定义转盘中抽取结果；点击“就喝这个”后，饮品信息进入记录草稿，购买后补充规格与价格即可保存。",
        scenario_numbering,
    )
    add_number(
        doc,
        "周末的生活回看",
        "：用户在“我的”查看一周饮品日历，再进入月历回看某一天的饮品与备注，重新发现当时的口味和心情。",
        scenario_numbering,
    )

    add_heading(doc, "体验闭环", 2)
    flow = doc.add_paragraph()
    flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    flow.paragraph_format.space_before = Pt(14)
    flow.paragraph_format.space_after = Pt(16)
    add_text(flow, "犹豫", size=12, bold=True, color=COFFEE)
    add_text(flow, "  →  ", size=12, color=CORAL)
    add_text(flow, "Choice One", size=12, bold=True, color=CORAL)
    add_text(flow, "  →  ", size=12, color=CORAL)
    add_text(flow, "预填记录", size=12, bold=True, color=COFFEE)
    add_text(flow, "  →  ", size=12, color=CORAL)
    add_text(flow, "用户确认", size=12, bold=True, color=COFFEE)
    add_text(flow, "  →  ", size=12, color=CORAL)
    add_text(flow, "日历沉淀", size=12, bold=True, color=COFFEE)

    add_callout(doc, "饮品不只是消费记录，也是一天里的小小注脚。")

    add_heading(doc, "日历回看，把零散记录连起来", 2)
    add_body(
        doc,
        "月历用轻量标记提示有记录的日期。选择日期后，可以查看当日饮品、杯数和已知热量合计。",
    )
    add_body(
        doc,
        "个人页的一周饮品日历提供更近距离的反馈：七天饮品章与周汇总，让用户无需进入完整月历，也能感知本周的饮品节奏。",
    )

    add_page_break(doc)

    add_heading(doc, "视觉与品牌体验", 1)
    add_body(
        doc,
        "WhatsDrink 采用“温暖饮品日记”视觉方向：奶油纸张底色承载轻微纹理，咖啡棕用于文字与主体信息，珊瑚橙用于 Choice One 和主要操作。圆角卡片、印章式日期与手账感插画，让记录过程更像在为一天留下温柔注脚。",
    )
    images = doc.add_paragraph()
    images.alignment = WD_ALIGN_PARAGRAPH.CENTER
    images.paragraph_format.space_before = Pt(8)
    add_image_with_alt(images, HOME_IMAGE, 2.36, "WhatsDrink 首页概念图，展示今日饮品时间线")
    add_text(images, "     ", size=8)
    add_image_with_alt(images, PROFILE_IMAGE, 2.36, "WhatsDrink 个人页概念图，展示一周饮品日历")
    caption = doc.add_paragraph(style="Caption WhatsDrink")
    add_text(caption, "首页概念图：今日饮品时间线                         个人页概念图：一周饮品日历", size=9, color=MUTED)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(2)
    add_text(note, "产品概念视觉 · 390 × 844 移动端方向", size=8.5, color=MUTED)

    add_page_break(doc)

    add_heading(doc, "产品特色", 1)
    add_bullet(doc, "记录与决策形成闭环：随机结果不是终点，而是下一条饮品记录的起点。")
    add_bullet(doc, "结构化目录与自由输入并存：常喝品牌可以快速选择，小众或自制饮品也不会被限制。")
    add_bullet(doc, "对不完整数据保持诚实：未知热量不参与合计，避免制造虚假的精确感。")
    add_bullet(doc, "私密优先：一期不设置动态、排行或社交压力，用户只为自己记录。")
    add_bullet(doc, "轻量但可回看：今日、周视图和月历分别覆盖即时记录、短期反馈和长期回顾。")

    add_heading(doc, "一期范围", 2)
    scope = doc.add_table(rows=1, cols=2)
    scope.style = "Table Grid"
    set_table_geometry(scope, [4680, 4680])
    mark_header_row(scope.rows[0])
    for idx, text in enumerate(("已包含", "暂不包含")):
        cell = scope.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT if idx == 0 else "F5F2EE")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, text, size=10, bold=True, color=COFFEE)
    in_scope = [
        "咖啡、奶茶两类饮品记录",
        "12 个品牌与 36 个代表单品",
        "今日、周视图与月历回看",
        "多转盘管理与等概率抽取",
        "资料、隐私与账号管理",
        "本地体验与草稿保留",
    ]
    out_scope = [
        "好友动态、排行榜与社区",
        "门店定位、外卖或支付跳转",
        "订阅付费与运营后台",
        "权重抽取与智能推荐",
        "健康、医疗或减脂建议",
        "品牌 Logo 与商业合作能力",
    ]
    row = scope.add_row()
    for idx, items in enumerate((in_scope, out_scope)):
        cell = row.cells[idx]
        set_cell_width(cell, 4680)
        set_cell_margins(cell, 120, 160, 120, 160)
        cell.text = ""
        for item in items:
            p = cell.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.18)
            p.paragraph_format.space_after = Pt(5)
            add_text(p, item, size=9.5)
        if cell.paragraphs and not cell.paragraphs[0].text:
            cell._tc.remove(cell.paragraphs[0]._p)

    add_page_break(doc)

    add_heading(doc, "技术与隐私", 1)
    add_body(
        doc,
        "小程序采用原生 TypeScript、WXML 与 WXSS 开发，后端使用微信云开发数据库、云函数与云存储。登录身份由云函数从微信上下文取得 OpenID，客户端不接触 session_key，也不提交可信用户 ID。",
    )
    add_body(
        doc,
        "个人资料、饮品记录、转盘与照片仅允许本人访问。记录和转盘写入由云函数鉴权；注销账号时级联删除个人资料、记录、转盘和上传照片。昵称与头像可跳过，照片权限仅在用户主动添加照片时请求。",
    )
    add_body(
        doc,
        "品牌目录采用版本化种子数据维护。无云环境或网络失败时，小程序保留本地体验与未提交草稿，不会把失败伪装成保存成功。",
    )

    add_heading(doc, "当前状态", 2)
    add_callout(doc, "一期已完成 8 个页面、5 个云函数、核心数据模型与 11 项自动化测试。")
    add_bullet(doc, "已覆盖记录新增、编辑、删除与汇总同步。")
    add_bullet(doc, "已覆盖日期日历、跨月逻辑、转盘边界与草稿映射。")
    add_bullet(doc, "接入真实 AppID 和云环境后，可在微信开发者工具中完成设备适配、弱网与隐私授权验收。")

    add_heading(doc, "后续方向", 2)
    add_body(
        doc,
        "未来可在不破坏“私密、轻量”定位的前提下，探索更多饮品类别、个人趋势总结、更丰富的记录回顾，以及由用户主动开启的个性化推荐。",
    )

    closing = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    closing.paragraph_format.space_before = Pt(28)
    closing.paragraph_format.space_after = Pt(4)
    add_text(closing, "WhatsDrink", size=15, bold=True, color=COFFEE)
    tagline = doc.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(tagline, "记录每一杯，也轻松决定下一杯。", size=11, color=CORAL)

    disclaimer = doc.add_paragraph()
    disclaimer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    disclaimer.paragraph_format.space_before = Pt(16)
    add_text(disclaimer, "参考热量仅用于日常记录，不构成医疗、营养或减脂建议。", size=8.5, color=MUTED)

    doc.core_properties.title = "WhatsDrink 微信小程序产品介绍"
    doc.core_properties.subject = "产品定位、核心体验、功能范围、技术与隐私说明"
    doc.core_properties.author = "WhatsDrink Product Team"
    doc.core_properties.keywords = "WhatsDrink, 微信小程序, 咖啡, 奶茶, Choice One, 饮品记录"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
